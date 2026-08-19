from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "VaultLab_Research_Paper.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_font(run, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic


def add_para(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, bold=True)
    r.font.size = Pt(14 if level == 1 else 12)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r)
    return p


def add_table(doc):
    add_heading(doc, "Table 1. Prototype variables and research role", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = [Inches(1.65), Inches(2.45), Inches(2.4)]
    headers = ["Variable", "How VaultLab estimates it", "Use in model"]
    for i, cell in enumerate(table.rows[0].cells):
        cell.text = headers[i]
        cell.width = widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell)
        for p in cell.paragraphs:
            for run in p.runs:
                set_font(run, bold=True)
    rows = [
        ("Approach speed", "Approach distance divided by approach time.", "Predictor of make/miss and performance potential."),
        ("Takeoff angle", "Pose landmarks near the estimated takeoff frame.", "Technique feature for takeoff posture and energy transfer."),
        ("Inversion angle", "Shoulder-hip-ankle angle near maximum inversion.", "Technique feature for swing and rock-back quality."),
        ("Grip over bar", "Grip height minus bar height, converted to meters.", "Context feature connecting technique demand to bar height."),
        ("Flight time", "Clearance timestamp minus takeoff timestamp.", "Phase timing feature for comparing attempts."),
        ("Result label", "Athlete records make or miss.", "Target variable for logistic regression."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].width = widths[i]
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_font(run)
                    run.font.size = Pt(10)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    styles["Normal"].font.size = Pt(12)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = 1.5
    r = title.add_run("VaultLab: AI-Powered Pole Vault Video Analysis for Technique Feedback and Performance Prediction")
    set_font(r, bold=True)
    r.font.size = Pt(14)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.line_spacing = 1.5
    for text in ["Jaya Xu", "\nSTEM Research & Innovation - Information Technology", "\nAugust 2026"]:
        run = meta.add_run(text)
        set_font(run)

    add_heading(doc, "Project Objectives")
    add_para(doc, "VaultLab is a web application that helps pole vault athletes convert ordinary side-view vault videos into measurable technical feedback. The project objective is to build an accessible AI-assisted analysis system that detects body landmarks, estimates vault phases, calculates technique variables, and prepares a dataset that can later train a predictive model for whether a vault is a make or miss.")
    add_para(doc, "The research question is: Can computer vision and interpretable machine learning identify which measurable technique variables are most associated with successful pole vault clearances?")
    add_para(doc, "The current prototype focuses on three deliverables: a React web app for video upload and attempt logging, a pose-estimation pipeline for extracting joint-based measurements, and a research workflow that exports labeled attempt data for logistic regression.")

    add_heading(doc, "Background and Motivation")
    add_para(doc, "Pole vault is a technically complex event because performance depends on runway speed, pole plant timing, takeoff mechanics, swing, inversion, and bar clearance. Athletes often rely on coaches visually reviewing video, but not every athlete has access to advanced biomechanical tools. A low-cost web app that extracts approximate measurements from a phone video could make feedback more available to developing athletes.")
    add_para(doc, "Biomechanics research supports the importance of measurable variables. Linthorne and colleagues reported that peak vault height increased by about 0.54 m for each 1 m/s increase in run-up velocity in a case study of an experienced vaulter. This supports using approach speed as a core feature rather than treating vault success as only a subjective technique judgment.")
    add_para(doc, "Recent markerless biomechanics studies show that computer vision is becoming realistic for pole vault analysis, but also show why a student project should be honest about uncertainty. Bissas and Cronin evaluated a markerless system in an elite competition setting using four cameras and reported strong agreement for center-of-mass and step velocity variables, while joint angles at takeoff remained less precise. This motivates VaultLab's design: show estimated values, preserve manual review, and collect data for validation.")

    add_heading(doc, "Methodology / Technical Approach")
    add_para(doc, "VaultLab uses a React + Vite + TypeScript frontend hosted on Firebase. The athlete signs in or enters competition demo mode, uploads a vault video, enters known attempt information such as bar height, grip height, approach distance, approach time, and make/miss result, then runs full-video pose analysis.")
    add_para(doc, "The prototype uses MediaPipe Pose Landmarker in the browser. According to Google's documentation, the task can detect human body landmarks from images or video and returns pose landmarks in image coordinates and world coordinates. VaultLab samples frames across the uploaded video, applies pose estimation, and draws a skeleton overlay so users can visually inspect whether the AI is tracking the athlete.")
    add_para(doc, "The app then estimates three major phases: takeoff, maximum inversion, and clearance. From these frames, it calculates or stores the following variables.")
    add_table(doc)
    add_para(doc, "The next research layer is logistic regression. Logistic regression is appropriate for this first model because the target variable is binary: make or miss. The exported dataset can be represented as rows of attempts and columns of features. A future Python notebook will standardize features, train a logistic regression classifier, evaluate accuracy with train/test splits or cross-validation, and inspect feature coefficients to explain which variables have the strongest relationship with success.")

    add_heading(doc, "Key Innovation")
    add_para(doc, "The key innovation is not only uploading a sports video. The innovation is connecting four layers into one athlete-facing workflow: pose estimation, vault-specific biomechanical feature extraction, interpretable predictive modeling, and coaching feedback. Most video tools allow manual drawing or review; VaultLab aims to turn video into a structured dataset that can support both athlete learning and research.")
    add_para(doc, "This is especially relevant because pole-vault-specific public datasets and automated phase-segmentation tools are limited. VaultLab therefore treats the app itself as a data collection instrument. Each saved attempt becomes a labeled example that can improve the model over time.")

    add_heading(doc, "Results / Findings")
    add_para(doc, "The current prototype has been implemented and deployed. It can authenticate users, offer a demo mode for judges, upload and preview a full vault video, run browser-based pose estimation, display joint markers on the video, estimate takeoff and inversion measurements, calculate approach speed and flight time when enough data is entered, store attempts in a user-specific log, and export attempt data as CSV.")
    add_para(doc, "The prototype also includes a Competition Research Mode panel. This panel counts total saved attempts, identifies how many rows are model-ready, displays the make rate, and clarifies that the prediction target is make/miss. This directly connects the web app to the planned logistic-regression study.")
    add_para(doc, "The most important early finding is a design finding: the app must present measurements as estimates, not absolute truth. Single-camera videos can have camera-angle error, motion blur, occlusion from the pole, and inverted poses that are difficult for general pose models. Therefore, VaultLab emphasizes reviewable phase estimates and dataset export before claiming final predictive accuracy.")

    add_heading(doc, "Impact / Future Development")
    add_para(doc, "VaultLab could help high school athletes and coaches receive more objective feedback from videos they already record. Instead of saying only that a vault looked good or bad, the app encourages athletes to compare speed, takeoff posture, inversion, grip, and result across attempts.")
    add_para(doc, "The next stage is to collect 30-50 labeled vault attempts, train the first logistic regression model, and report metrics such as accuracy, precision, recall, and the most influential coefficients. Later versions could add cloud storage, Firestore data sync, a Python backend with YOLO Pose or RTMPose, pole and bar detection, and two-camera calibration for more accurate measurements.")
    add_para(doc, "For the competition, the project demonstrates a full STEM innovation cycle: identify a real athletic problem, build a working AI prototype, connect it to a research question, collect structured data, and design an interpretable model that can guide improvement.")

    add_heading(doc, "References")
    refs = [
        "Bissas, A., & Cronin, N. J. (2026). Evaluating markerless biomechanical analysis in a real-world pole vault competition setting. PLOS One, 21(3), e0329987. https://doi.org/10.1371/journal.pone.0329987",
        "Google AI Edge. (2026). Pose landmark detection guide. https://developers.google.com/edge/mediapipe/solutions/vision/pose_landmarker",
        "Google Firebase. (2026). Get started with Firebase Authentication on websites. https://firebase.google.com/docs/auth/web/start",
        "Linthorne, N. P., Guzman, M. S., & Bridgett, L. A. (2012). Effects of run-up velocity on performance, kinematics, and energy exchanges in the pole vault. Journal of Sports Science and Medicine, 11(2), 245-254. https://pmc.ncbi.nlm.nih.gov/articles/PMC3737876/",
        "scikit-learn developers. (2026). LogisticRegression. https://scikit-learn.org/stable/modules/generated/sklearn.linear_model.LogisticRegression.html",
    ]
    for ref in refs:
        add_para(doc, ref)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
